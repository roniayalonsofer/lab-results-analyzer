"""Regression tests for Machon HaNeft (מכון הנפט) certificates.

Uses small synthetic workbooks that mimic the real layout, so no client data
is committed. Run:  py -3 -m pytest tests -q   (or: py -3 tests/test_machon_haneft.py)
"""
import io
import os
import sys

import openpyxl

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "soil_lab_tool"))

from parsers import auto_detect_lab, get_parser  # noqa: E402
from parsers.soil.machon_haneft import MachonHaneftSoilParser  # noqa: E402

SAMPLES = ["ק-1 -3.0", "ק-1 -6.0", "ק-1 -10.5", "ק-1 -13.5", "ק-10 -3.0", "ק-10 -6.0"]


def _workbook() -> bytes:
    """Sheet 'TPH' (the name Alchem's detector also matches) holding a TPH block
    and, below it, a BTEX block: header, units, LOD, LOQ, method, CAS, samples."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TPH"
    ws["A8"] = "תעודת בדיקה מס' 1/2026"
    ws["A12"] = "תאריך לקיחת המדגם:"
    ws["C29"], ws["D29"], ws["E29"] = "TPH", "DRO", "ORO"
    ws["A30"] = "יחידות"; ws["C30"] = "mg/kg"
    ws["A31"] = "שיטה"; ws["C31"] = "EPA 8015D"
    ws["A32"] = "LOQ"; ws["C32"] = 50
    for i, s in enumerate(SAMPLES):
        ws.cell(row=33 + i, column=1, value=s)
        ws.cell(row=33 + i, column=3, value=100 + i)
        ws.cell(row=33 + i, column=4, value=60 + i)
        ws.cell(row=33 + i, column=5, value=40 + i)
    hdr = 50
    for ci, h in enumerate(["MTBE", "בנזן", "טולואן", "אתיל-בנזן", "קסילן"], start=3):
        ws.cell(row=hdr, column=ci, value=h)
    ws.cell(row=hdr, column=1, value="בדיקה")
    ws.cell(row=hdr + 1, column=1, value="יחידות")
    ws.cell(row=hdr + 2, column=1, value="גבול גילוי הבדיקה")
    ws.cell(row=hdr + 3, column=1, value="גבול כימות הבדיקה")
    ws.cell(row=hdr + 4, column=1, value="שיטה"); ws.cell(row=hdr + 4, column=3, value="EPA 8260D")
    for ci, cas in enumerate(["1634-04-4", "71-43-2", "108-88-3", "100-41-4", "1330-20-7"], start=3):
        ws.cell(row=hdr + 5, column=ci, value=f"CAS No.{cas}")
        ws.cell(row=hdr + 2, column=ci, value=0.01)
        ws.cell(row=hdr + 3, column=ci, value=0.03)
    for i, s in enumerate(SAMPLES):
        ws.cell(row=hdr + 6 + i, column=1, value=s)
        for ci in range(3, 8):
            ws.cell(row=hdr + 6 + i, column=ci, value=1.5 + i)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_btex_block_keeps_first_samples():
    """Regression: '' in _SKIP_PREFIXES made every row look like metadata, so the
    first samples of the BTEX table were silently dropped."""
    recs = MachonHaneftSoilParser().parse(io.BytesIO(_workbook()))
    btex = [r for r in recs if r["analysis_type"] == "SOIL_MBTEX"]
    assert {r["sample_id"] for r in btex} == set(SAMPLES)
    assert len(btex) == len(SAMPLES) * 5
    first = next(r for r in btex if r["sample_id"] == SAMPLES[0] and r["compound"] == "בנזן")
    assert first["value"] == 1.5


def test_tph_and_btex_both_present():
    recs = get_parser("מכון הנפט", "soil").parse(io.BytesIO(_workbook()))
    types = {r["analysis_type"] for r in recs}
    assert {"SOIL_TPH", "SOIL_MBTEX"} <= types


def test_excel_detected_as_machon_haneft_not_alchem():
    """Regression: a first sheet named 'TPH' used to be claimed by Alchem."""
    assert auto_detect_lab("report.xlsx", _workbook()) == "מכון הנפט"


if __name__ == "__main__":
    for name, fn in list(globals().items()):
        if name.startswith("test_"):
            fn()
            print("ok", name)
